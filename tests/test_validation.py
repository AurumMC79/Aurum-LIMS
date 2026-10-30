"""Validation package: traceability, installation checks, the generated documents and change control."""

import json

import pytest
from docx import Document
from sqlalchemy import select, text

from aurum_lims.db.models import AuditEntry
from aurum_lims.errors import PermissionDenied
from aurum_lims.security.roles import Role
from aurum_lims.services import backup
from aurum_lims.services.audit import AuditAction
from aurum_lims.ui.views.settings import SettingsView
from aurum_lims.validation import content, package


def test_every_requirement_is_traced_to_a_test():
    ids = content.requirement_ids()
    assert len(ids) == len(set(ids)), "requirement IDs must be unique"
    tests = content.all_tests()
    assert len({t.id for t in tests}) == len(tests), "test IDs must be unique"
    for test in tests:
        assert test.steps and test.requirements, test.id
        assert set(test.requirements) <= set(ids), f"{test.id} refers to an unknown requirement"
    covered = {r for t in tests for r in t.requirements}
    assert covered == set(ids), f"not tested: {sorted(set(ids) - covered)}"
    high = [r.id for r in content.REQUIREMENTS if r.risk == content.HIGH]
    assert all(row[3] for row in package.traceability_rows() if row[0] in high), "High risk needs an OQ test"


def test_installation_checks(db, admin):
    backup.backup_now(db, admin)
    with db.session() as s:
        checks = package.installation_checks(db, s)
    assert all(c.passed for c in checks), [(c.name, c.observed) for c in checks if not c.passed]
    # A missing protection trigger is caught.
    with db.engine.begin() as conn:
        conn.execute(text("DROP TRIGGER certificates_no_update"))
    with db.session() as s:
        failed = {c.name: c.observed for c in package.installation_checks(db, s) if not c.passed}
    assert failed == {"Immutable records are protected in the database": "missing: certificates_no_update"}


def test_backup_check_fails_without_a_recent_backup(db, admin):
    with db.session() as s:
        failed = [c.name for c in package.installation_checks(db, s) if not c.passed]
    assert failed == ["Automatic backups are on and recent"]


def test_generate_package(db, admin, make_user, tmp_path):
    with pytest.raises(PermissionDenied):
        package.generate(db, make_user("ann", Role.ANALYST), tmp_path)
    qa = make_user("qa1", Role.QA_REVIEWER)  # QA can create it as well as administrators
    backup.backup_now(db, admin)
    result = package.generate(db, qa, tmp_path, "VAL-LAB")
    names = sorted(p.name for p in result.folder.iterdir())
    assert names == ["01 Validation Plan.docx", "02 User Requirements Specification.docx",
                     "03 Installation Qualification Protocol.docx", "04 Operational Qualification Protocol.docx",
                     "05 Performance Qualification Protocol.docx", "06 Traceability Matrix.docx",
                     "07 Validation Summary Report.docx", "program_file_manifest.csv", "traceability_matrix.csv"]
    assert result.checks_passed

    iq = Document(result.folder / "03 Installation Qualification Protocol.docx")
    body = "\n".join(p.text for p in iq.paragraphs)
    assert result.fingerprint in body
    cells = [c.text for t in iq.tables for row in t.rows for c in row.cells]
    assert "Database integrity (SQLite integrity_check)" in cells and "VAL-LAB-IQ-001" in cells
    assert "aurum_lims" not in iq.sections[0].footer.paragraphs[0].text  # footer is the document number
    assert "VAL-LAB-IQ-001" in iq.sections[0].footer.paragraphs[0].text

    oq = Document(result.folder / "04 Operational Qualification Protocol.docx")
    oq_cells = {c.text for t in oq.tables for row in t.rows for c in row.cells}
    assert all(t.id in oq_cells for a in content.OQ for t in a.tests)
    urs = Document(result.folder / "02 User Requirements Specification.docx")
    urs_cells = {c.text for t in urs.tables for row in t.rows for c in row.cells}
    assert set(content.requirement_ids()) <= urs_cells

    manifest = (result.folder / "program_file_manifest.csv").read_text(encoding="utf-8-sig")
    assert "services/backup.py" in manifest and result.fingerprint in manifest
    with db.session() as s:
        entry = s.scalars(select(AuditEntry).where(AuditEntry.action == AuditAction.VALIDATION_PACKAGE_CREATED)).one()
        assert entry.username == "qa1" and json.loads(entry.new_values)["program_fingerprint"] == result.fingerprint
        assert package.last_package(s).id == entry.id


def test_fingerprint_changes_when_a_program_file_changes(tmp_path, monkeypatch):
    (tmp_path / "app.py").write_text("print('v1')")
    (tmp_path / "notes.txt").write_text("not a program file")
    monkeypatch.setattr(package, "program_root", lambda: tmp_path)
    first = package.program_manifest()
    assert [e.path for e in first] == ["app.py"]
    (tmp_path / "notes.txt").write_text("changed")
    assert package.manifest_fingerprint(package.program_manifest()) == package.manifest_fingerprint(first)
    (tmp_path / "app.py").write_text("print('v2')")
    assert package.manifest_fingerprint(package.program_manifest()) != package.manifest_fingerprint(first)


def test_validation_tab(qtbot, db, admin, make_user, tmp_path, monkeypatch):
    view = SettingsView(db, admin)
    qtbot.addWidget(view)
    view.refresh()
    tab = view.validation_tab
    assert "No validation package" in tab.status.text()
    backup.backup_now(db, admin)
    tab.folder.setText(str(tmp_path))
    monkeypatch.setattr("PySide6.QtWidgets.QMessageBox.information", lambda *a, **k: None)
    tab._create()
    assert "All 9 automatic installation checks passed" in tab.result.text()
    assert "unchanged since" in tab.status.text() and tab.open_button.isEnabled()
    analyst_view = SettingsView(db, make_user("ann", Role.ANALYST))
    qtbot.addWidget(analyst_view)
    assert analyst_view.validation_tab is None
