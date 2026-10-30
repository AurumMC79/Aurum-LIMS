"""Generates the validation package as editable Word documents.

The package follows the usual risk-based (GAMP 5) life cycle for a laboratory
system: validation plan → user requirements (URS) → installation (IQ),
operational (OQ) and performance (PQ) qualification → traceability matrix →
summary report. Everything except the IQ facts is drawn from
:mod:`aurum_lims.validation.content`.

The IQ is filled in with this installation's real details at the moment of
generation: versions, the database and its schema revision, automatic checks
(database integrity, audit-trail chain, protection triggers, backups) and a
SHA-256 manifest of the program files. The manifest's overall fingerprint
identifies exactly which software was qualified; generating the package again
after an update shows whether the program changed.
"""

from __future__ import annotations

import csv
import hashlib
import platform
import sqlite3
import sys
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from importlib import metadata
from pathlib import Path

from alembic.config import Config
from alembic.script import ScriptDirectory
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy import select, text

import aurum_lims
from aurum_lims import config
from aurum_lims.db.models import AuditEntry, User
from aurum_lims.db.session import MIGRATIONS_DIR, Database
from aurum_lims.security.identity import SessionUser
from aurum_lims.security.roles import ROLE_LABELS, Permission, require_any
from aurum_lims.services import audit, backup, lab_info
from aurum_lims.services.audit import AuditAction
from aurum_lims.services.settings import get_security_settings
from aurum_lims.validation.content import OQ, PQ, REQUIREMENTS, TestCase, all_tests

GOLD = RGBColor(0x9A, 0x7B, 0x1C)
HEADER_FILL = "F3ECD6"
PROTECTED_TABLES = ["audit_entries", "electronic_signatures", "custody_events", "reportable_results",
                    "sample_decisions", "certificates", "instrument_events", "lot_transactions",
                    "document_revisions", "imported_files"]
PROGRAM_SUFFIXES = {".py", ".qss", ".ico", ".mako", ".exe", ".dll", ".pyd", ".zip"}
LIBRARIES = ["PySide6-Essentials", "SQLAlchemy", "alembic", "argon2-cffi", "python-docx"]

DOCUMENTS = [
    ("01", "Validation Plan", "VP"), ("02", "User Requirements Specification", "URS"),
    ("03", "Installation Qualification Protocol", "IQ"), ("04", "Operational Qualification Protocol", "OQ"),
    ("05", "Performance Qualification Protocol", "PQ"), ("06", "Traceability Matrix", "TM"),
    ("07", "Validation Summary Report", "VSR"),
]


# ---------------------------------------------------------------------------
# facts about this installation
# ---------------------------------------------------------------------------


@dataclass
class Check:
    name: str
    expected: str
    observed: str
    passed: bool


@dataclass
class ManifestEntry:
    path: str
    size: int
    sha256: str


def program_root() -> Path:
    """Where the program files are: the PyInstaller folder, or the source package."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).parent
    return Path(aurum_lims.__file__).parent


def program_manifest() -> list[ManifestEntry]:
    root = program_root()
    entries = []
    for path in sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in PROGRAM_SUFFIXES
                       and "__pycache__" not in p.parts):
        data = path.read_bytes()
        entries.append(ManifestEntry(path.relative_to(root).as_posix(), len(data), hashlib.sha256(data).hexdigest()))
    return entries


def manifest_fingerprint(entries: list[ManifestEntry]) -> str:
    """One SHA-256 over every file's path and hash: changes if any program file changes."""
    digest = hashlib.sha256()
    for e in entries:
        digest.update(f"{e.path}\t{e.sha256}\n".encode())
    return digest.hexdigest()


def _size(size: int) -> str:
    return f"{size / 1_048_576:.2f} MB" if size >= 1_048_576 else f"{size / 1024:.0f} KB"


def _version(package: str) -> str:
    try:
        return metadata.version(package)
    except metadata.PackageNotFoundError:
        return "not installed"


def _head_revision() -> str:
    cfg = Config()
    cfg.set_main_option("script_location", str(MIGRATIONS_DIR))
    return ScriptDirectory.from_config(cfg).get_current_head()


def installation_facts(db: Database) -> list[tuple[str, str]]:
    from PySide6 import __version__ as pyside_version
    from PySide6.QtCore import qVersion

    db_path = backup._sqlite_path(db)
    with db.engine.connect() as conn:
        revision = conn.exec_driver_sql("SELECT version_num FROM alembic_version").scalar()
        journal = conn.exec_driver_sql("PRAGMA journal_mode").scalar()
    facts = [
        ("Software", f"{config.APP_NAME} {config.APP_VERSION}"),
        ("Program location", str(program_root())),
        ("Run as", "Packaged application (AurumLIMS.exe)" if getattr(sys, "frozen", False) else
         f"Python source ({sys.executable})"),
        ("Computer name", platform.node()),
        ("Operating system", platform.platform()),
        ("Python", platform.python_version()),
        ("Qt / PySide6", f"Qt {qVersion()} / PySide6 {pyside_version}"),
        *[(f"Library: {name}", _version(name)) for name in LIBRARIES if name != "PySide6-Essentials"],
        ("Data folder", str(config.data_dir())),
        ("Database", str(db_path) if db_path else db.url),
        ("Database size", _size(db_path.stat().st_size) if db_path and db_path.exists() else "—"),
        ("Database schema revision", f"{revision} (latest for this version: {_head_revision()})"),
        ("Database journal mode", str(journal).upper()),
        ("Log folder", str(config.log_dir())),
    ]
    return facts


def installation_checks(db: Database, session) -> list[Check]:
    checks = []
    db_path = backup._sqlite_path(db)
    if db_path is not None and db_path.exists():
        conn = sqlite3.connect(f"file:{db_path.as_posix()}?mode=ro", uri=True)
        try:
            result = conn.execute("PRAGMA integrity_check").fetchone()[0]
            triggers = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type = 'trigger'")}
        finally:
            conn.close()
        checks.append(Check("Database integrity (SQLite integrity_check)", "ok", result, result == "ok"))
    else:
        triggers = {r[0] for r in session.execute(text("SELECT name FROM sqlite_master WHERE type = 'trigger'"))}
    revision = session.execute(text("SELECT version_num FROM alembic_version")).scalar()
    head = _head_revision()
    checks.append(Check("Database schema is up to date", head, revision, revision == head))
    report = audit.verify_integrity(session, record_result=False)
    checks.append(Check("Audit trail hash chain verified", "intact",
                        f"{report.entries_checked} entries, " + ("intact" if report.ok else
                                                                 f"{len(report.problems)} problem(s)"), report.ok))
    missing = [f"{t}_no_{op}" for t in PROTECTED_TABLES for op in ("update", "delete") if f"{t}_no_{op}" not in triggers]
    checks.append(Check("Immutable records are protected in the database",
                        f"{len(PROTECTED_TABLES) * 2} protection triggers", "all present" if not missing else
                        "missing: " + ", ".join(missing), not missing))
    checks.append(Check("Samples cannot be deleted", "samples_no_delete trigger",
                        "present" if "samples_no_delete" in triggers else "missing", "samples_no_delete" in triggers))
    settings = backup.get_settings(session)
    last = backup.last_backup_time(session)
    folder = settings.path
    writable = False
    try:
        folder.mkdir(parents=True, exist_ok=True)
        probe = folder / ".aurum-write-test"
        probe.write_text("ok")
        probe.unlink()
        writable = True
    except OSError:
        pass
    checks.append(Check("Backup folder is writable", str(folder), "writable" if writable else "NOT writable", writable))
    due = settings.interval_hours and last and datetime.now(timezone.utc) - last < timedelta(
        hours=settings.interval_hours * 2)
    checks.append(Check("Automatic backups are on and recent",
                        f"every {settings.interval_hours} h" if settings.interval_hours else "on",
                        f"last backup {last.astimezone():%Y-%m-%d %H:%M}" if last else "no backup yet", bool(due)))
    admins = session.scalar(select(User.id).where(User.role == "master_administrator", User.is_active.is_(True)))
    checks.append(Check("An active Master Administrator exists", "yes", "yes" if admins else "no", admins is not None))
    ok_python = sys.version_info >= (3, 11)
    checks.append(Check("Python version supported", "3.11 or later", platform.python_version(), ok_python))
    return checks


# ---------------------------------------------------------------------------
# Word helpers
# ---------------------------------------------------------------------------


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _field(paragraph, instruction: str) -> None:
    """Insert a Word field (PAGE, NUMPAGES) that Word fills in when the document is opened."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
            run._r.append(element)
        elif text == instruction:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = f" {instruction} "
            run._r.append(element)
        else:
            run._r.append(OxmlElement("w:t"))
            run._r[-1].text = text


class Doc:
    """A validation document with the common header, footer, title block and approvals."""

    def __init__(self, code: str, title: str, number: str, lab: str, landscape: bool = False):
        self.document = Document()
        self.code, self.title, self.number = code, title, number
        styles = self.document.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(10)
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
            styles[name].font.name = "Calibri"
            styles[name].font.size = Pt(size)
            styles[name].font.color.rgb = GOLD
        section = self.document.sections[0]
        if landscape:
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2))
        section.top_margin = section.bottom_margin = Cm(1.8)
        header = section.header.paragraphs[0]
        header.text = f"{lab}\t{title}"
        header.runs[0].font.size = Pt(8)
        footer = section.footer.paragraphs[0]
        footer.add_run(f"{number}  ·  Version 0.1 (draft)  ·  Page ").font.size = Pt(8)
        _field(footer, "PAGE")
        footer.add_run(" of ").font.size = Pt(8)
        _field(footer, "NUMPAGES")
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

        heading = self.document.add_paragraph()
        run = heading.add_run(title)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(20), GOLD
        self.document.add_paragraph(f"{config.APP_NAME} {config.APP_VERSION}  ·  {lab}")
        self.table(["Document number", "Version", "Status", "Prepared on"],
                   [[number, "0.1", "Draft – for review", date.today().isoformat()]])

    # ----- building blocks -----
    def h1(self, text: str):
        return self.document.add_heading(text, level=1)

    def h2(self, text: str):
        return self.document.add_heading(text, level=2)

    def p(self, text: str = "", bold: bool = False, italic: bool = False, size: int | None = None):
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        return paragraph

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.document.add_paragraph(item, style="List Bullet")

    def table(self, headers: list[str], rows: list[list], widths: list[float] | None = None,
              font: int = 9, blank_rows: int = 0):
        table = self.document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = ""
            run = cell.paragraphs[0].add_run(text)
            run.bold, run.font.size = True, Pt(font)
            _shade(cell, HEADER_FILL)
        for row in [*rows, *[[""] * len(headers)] * blank_rows]:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = ""
                run = cell.paragraphs[0].add_run("" if value is None else str(value))
                run.font.size = Pt(font)
        if widths:
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        # Repeat the heading row on every page.
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_properties.append(repeat)
        self.document.add_paragraph()
        return table

    def approvals(self, title: str = "Approvals", roles=("Prepared by", "Reviewed by", "Approved by (QA)")) -> None:
        self.h2(title)
        self.table(["Role", "Name", "Signature", "Date"], [[r, "", "", ""] for r in roles], [4, 5, 5, 3])

    def history(self) -> None:
        self.h2("Revision history")
        self.table(["Version", "Date", "Description of change", "Author"],
                   [["0.1", date.today().isoformat(), f"Generated by {config.APP_NAME} {config.APP_VERSION}", ""]],
                   blank_rows=2)

    def page_break(self) -> None:
        self.document.add_page_break()

    def save(self, folder: Path, prefix: str) -> Path:
        path = folder / f"{prefix} {self.title}.docx"
        self.document.save(path)
        return path


def _test_case(doc: Doc, test: TestCase) -> None:
    doc.h2(f"{test.id}  {test.title}")
    doc.table(["Objective", "Requirements tested", "Prerequisites"],
              [[test.objective, ", ".join(test.requirements), test.prerequisites]], [6, 4, 7])
    doc.table(["Step", "Action", "Expected result", "Actual result", "Pass / Fail", "Initials / date"],
              [[str(i), action, expected, "", "", ""] for i, (action, expected) in enumerate(test.steps, start=1)],
              [1, 6.5, 6.5, 5.5, 1.8, 2.2], font=8)
    doc.table(["Test result", "Executed by / date", "Reviewed by / date", "Deviation no. (if any)"],
              [["☐ Pass   ☐ Fail", "", "", ""]], [4, 5.5, 5.5, 4])


# ---------------------------------------------------------------------------
# the documents
# ---------------------------------------------------------------------------


def _numbers(prefix: str) -> dict[str, str]:
    return {code: f"{prefix}-{code}-001" for _, _, code in DOCUMENTS}


def validation_plan(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("VP", "Validation Plan", numbers["VP"], lab)
    doc.approvals()
    doc.history()
    doc.h1("1  Purpose")
    doc.p(f"This plan describes how {lab} validates {config.APP_NAME} {config.APP_VERSION} for its intended use as "
          "the laboratory information management system, so that it consistently does what the laboratory needs and "
          "meets data-integrity expectations (21 CFR Part 11, EU GMP Annex 11, ISO/IEC 17025, ALCOA+).")
    doc.h1("2  Scope and system description")
    doc.p(f"{config.APP_NAME} is a Windows desktop application (Python, Qt) storing its data in a local SQLite "
          "database. It covers user accounts and roles, an audit trail and electronic signatures, samples and chain "
          "of custody, methods and specifications, result entry and calculation, instrument data import, OOS "
          "investigations, second-person review, QA decisions, certificates, instruments and maintenance, inventory "
          "and suppliers, controlled documents, reports and backups.")
    doc.p("Out of scope: the Windows operating system, network storage and the instruments' own software, which are "
          "qualified under their own procedures. Interfaces are limited to file import (CSV / TXT) and export (PDF, "
          "CSV).")
    doc.h1("3  Validation approach")
    doc.p("Risk-based approach following GAMP 5. The software is configured and was developed for the laboratory "
          "(GAMP software category 5, with category 4 configuration such as methods, specifications and settings). "
          "Effort is focused on requirements rated High GxP risk in the URS.")
    doc.table(["Deliverable", "Document number", "Purpose"], [
        ["Validation Plan", numbers["VP"], "This document: approach, responsibilities, acceptance criteria."],
        ["User Requirements Specification", numbers["URS"], "What the laboratory requires, with GxP risk."],
        ["Installation Qualification", numbers["IQ"], "The software is installed correctly in its environment."],
        ["Operational Qualification", numbers["OQ"], "Each function works as specified (test scripts)."],
        ["Performance Qualification", numbers["PQ"], "The system supports the laboratory's real work end to end."],
        ["Traceability Matrix", numbers["TM"], "Every requirement is tested."],
        ["Validation Summary Report", numbers["VSR"], "Results, deviations and the release decision."],
    ], [5, 4, 8])
    doc.p("Supplier testing: the software is developed with an automated regression test suite (covering security, "
          "audit trail, calculations, workflows and database upgrades) that is run on every change. Its results may "
          "be referenced as supporting evidence but do not replace the OQ.")
    doc.h1("4  Responsibilities")
    doc.table(["Role", "Name", "Responsibilities"], [
        ["System owner", "", "Owns the system and this validation; approves the URS and the summary report."],
        ["Quality Assurance", "", "Reviews and approves all protocols, deviations and the release decision."],
        ["IT / administrator", "", "Installs the software, executes the IQ, maintains backups and accounts."],
        ["Test executors", "", "Execute OQ / PQ tests and record results under good documentation practice."],
    ], [4, 4, 9])
    doc.h1("5  Acceptance criteria")
    doc.bullets(["All IQ checks pass, or deviations are resolved.",
                 "All OQ and PQ test steps meet their expected results, or deviations are assessed by QA as not "
                 "affecting the intended use.",
                 "Every URS requirement is traced to at least one passed test.",
                 "All High-risk deviations are closed before release."])
    doc.h1("6  Deviations")
    doc.p("Any result that does not match the expected result is recorded as a deviation (number, description, "
          "impact assessment, corrective action, re-test, QA approval) and listed in the summary report.")
    doc.h1("7  Change control and periodic review")
    doc.bullets([
        "Every software update is assessed under change control. The IQ is repeated (the program fingerprint shows "
        "whether the program changed) and affected OQ tests are re-executed.",
        "A pre-update database backup is made automatically before every database upgrade.",
        "The validated state, audit-trail reviews, user access and backups are reviewed at least annually.",
        "Users are trained on the procedures and the software before they are given an account."])
    return doc


def urs(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("URS", "User Requirements Specification", numbers["URS"], lab)
    doc.approvals()
    doc.history()
    doc.h1("1  Introduction")
    doc.p(f"These are {lab}'s requirements for its laboratory information management system. Each requirement has a "
          "unique ID used in the protocols and the traceability matrix, and a GxP risk rating: High = a failure could "
          "affect product quality, patient safety or data integrity; Medium = could affect compliance or efficiency; "
          "Low = convenience. Edit, add or remove requirements to fit the laboratory before approval.")
    areas: dict[str, list] = {}
    for r in REQUIREMENTS:
        areas.setdefault(r.area, []).append(r)
    for index, (area, items) in enumerate(areas.items(), start=2):
        doc.h1(f"{index}  {area}")
        doc.table(["ID", "Requirement", "GxP risk", "Supports"], [[r.id, r.text, r.risk, r.reference] for r in items],
                  [2.6, 9.5, 1.6, 3.3])
    doc.h1(f"{len(areas) + 2}  Roles")
    doc.p("The system provides these roles; their permissions are defined in the software and verified in the OQ.")
    doc.bullets(list(ROLE_LABELS.values()))
    return doc


def iq(lab: str, numbers: dict[str, str], facts, checks, manifest: list[ManifestEntry], security) -> Doc:
    doc = Doc("IQ", "Installation Qualification Protocol", numbers["IQ"], lab)
    doc.approvals("Protocol approval")
    doc.h1("1  Purpose")
    doc.p("To verify and document that the software is installed correctly on this computer, with the intended "
          "versions, a healthy database and its data-protection measures in place. The values below were recorded "
          f"automatically by {config.APP_NAME} on {datetime.now().astimezone():%Y-%m-%d %H:%M %Z}; the executor "
          "confirms each one.")
    doc.h1("2  Installation details")
    doc.table(["Item", "Recorded value", "Verified (initials / date)"], [[k, v, ""] for k, v in facts], [4.5, 9, 3.5])
    doc.h1("3  Automatic installation checks")
    doc.table(["Check", "Expected", "Observed", "Result"],
              [[c.name, c.expected, c.observed, "Pass" if c.passed else "FAIL"] for c in checks], [6, 4, 5, 2])
    doc.h1("4  Security configuration")
    doc.table(["Setting", "Configured value", "Approved value", "Initials / date"], [
        ["Lock screen after inactivity", f"{security.session_timeout_minutes} min", "", ""],
        ["Failed sign-ins before lockout", str(security.max_failed_logins), "", ""],
        ["Minimum password length", str(security.password_min_length), "", ""],
        ["Password expiry", f"{security.password_expiry_days} days" if security.password_expiry_days else "never", "",
         ""],
        ["Remembered previous passwords", str(security.password_history_count), "", ""],
    ], [6, 3.5, 3.5, 4])
    doc.h1("5  Manual installation checks")
    manual = [
        ("The desktop shortcut starts Aurum LIMS and the sign-in window appears.", "Starts; sign-in shown."),
        ("The computer clock is synchronised with a reliable time source (e.g. domain / NTP).", "Synchronised."),
        ("Windows access to the data folder is limited to authorised users.", "Only authorised accounts."),
        ("The backup folder is on a separate disk or network location covered by the IT backup.", "Confirmed."),
        ("Settings → Backups → Back up now, then Check selected backup.", "Backup made; integrity check passed."),
        ("Users: only named accounts exist; no shared or test accounts remain active.", "Confirmed."),
        ("Settings → Laboratory details are correct (printed on certificates).", "Correct."),
    ]
    doc.table(["#", "Check", "Expected", "Observed", "Pass / Fail", "Initials / date"],
              [[str(i), c, e, "", "", ""] for i, (c, e) in enumerate(manual, start=1)], [0.8, 6.5, 3.5, 3, 1.6, 2.2])
    doc.h1("6  Program files")
    doc.p(f"{len(manifest)} program files were fingerprinted (SHA-256). Overall program fingerprint:", bold=True)
    doc.p(manifest_fingerprint(manifest), size=9)
    doc.p("Record this fingerprint. If the package is generated again and the fingerprint differs, the program has "
          "changed and must go through change control. The full list is in Appendix A and in the file "
          "program_file_manifest.csv.")
    doc.h1("7  Deviations and conclusion")
    doc.table(["Deviation no.", "Description", "Resolution"], [], [3, 8, 6], blank_rows=3)
    doc.p("☐ The installation is qualified.   ☐ The installation is not qualified (see deviations).")
    doc.approvals("Execution sign-off", ("Executed by", "Reviewed by", "Approved by (QA)"))
    doc.page_break()
    doc.h1("Appendix A  Program file manifest")
    doc.table(["File", "Size (bytes)", "SHA-256"], [[e.path, f"{e.size:,}", e.sha256] for e in manifest],
              [6, 2.2, 8.8], font=7)
    return doc


def oq(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("OQ", "Operational Qualification Protocol", numbers["OQ"], lab, landscape=True)
    doc.approvals("Protocol approval")
    doc.h1("1  Purpose and instructions")
    doc.p("To verify that each function of the software works as required by the URS. Execute the tests on the "
          "qualified installation (see IQ) using dedicated test accounts and clearly marked test samples, or on a "
          "copy of the system. Steps marked IT use a copy of the database and are never performed on production "
          "data.")
    doc.bullets(["Record actual results in ink (or electronically under your document procedure); do not leave "
                 "blank fields – use N/A with initials.",
                 "Attach screenshots or printouts as evidence where useful and reference them in Actual result.",
                 "Any result that differs from the expected result is a deviation: stop, record it and inform QA.",
                 "Test accounts created for the OQ are deactivated after execution."])
    doc.h2("Test environment")
    doc.table(["Item", "Value"], [["Computer / installation (IQ reference)", ""], ["Software version",
                                                                                   config.APP_VERSION],
                                  ["Test accounts used", ""], ["Execution dates", ""]], [8, 15])
    for number, area in enumerate(OQ, start=2):
        doc.page_break()
        doc.h1(f"{number}  {area.name}")
        for test in area.tests:
            _test_case(doc, test)
    doc.page_break()
    doc.h1(f"{len(OQ) + 2}  Summary of results")
    doc.table(["Test", "Title", "Result", "Deviation no."],
              [[t.id, t.title, "", ""] for area in OQ for t in area.tests], [3, 12, 3, 4])
    doc.approvals("Execution sign-off", ("Executed by", "Reviewed by", "Approved by (QA)"))
    return doc


def pq(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("PQ", "Performance Qualification Protocol", numbers["PQ"], lab, landscape=True)
    doc.approvals("Protocol approval")
    doc.h1("1  Purpose and instructions")
    doc.p("To show that the system supports the laboratory's real work, with its own methods, specifications, "
          "instruments, users and procedures, in the production environment. Replace the generic steps with the "
          "laboratory's SOP steps and products before approving this protocol.")
    doc.h1("2  Scenarios")
    for test in PQ:
        _test_case(doc, test)
    doc.h1("3  Summary of results")
    doc.table(["Scenario", "Title", "Result", "Deviation no."], [[t.id, t.title, "", ""] for t in PQ], [3, 12, 3, 4])
    doc.approvals("Execution sign-off", ("Executed by", "Reviewed by", "Approved by (QA)"))
    return doc


def traceability_rows() -> list[list[str]]:
    tests = all_tests()
    rows = []
    for r in REQUIREMENTS:
        oq_ids = [t.id for t in tests if r.id in t.requirements and t.id.startswith("OQ")]
        pq_ids = [t.id for t in tests if r.id in t.requirements and t.id.startswith("PQ")]
        rows.append([r.id, r.text, r.risk, ", ".join(oq_ids), ", ".join(pq_ids)])
    return rows


def traceability(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("TM", "Traceability Matrix", numbers["TM"], lab, landscape=True)
    doc.approvals()
    doc.p("Each requirement of the URS and the tests that verify it. Complete the Result column from the executed "
          "protocols; every requirement must be covered by at least one passed test.")
    doc.table(["Requirement", "Description", "GxP risk", "OQ tests", "PQ tests", "Result"],
              [[*row, ""] for row in traceability_rows()], [2.6, 11, 1.6, 4.2, 2.6, 2], font=8)
    return doc


def summary_report(lab: str, numbers: dict[str, str]) -> Doc:
    doc = Doc("VSR", "Validation Summary Report", numbers["VSR"], lab)
    doc.approvals()
    doc.history()
    doc.h1("1  Summary")
    doc.p(f"[Summarise the validation of {config.APP_NAME} {config.APP_VERSION} carried out under "
          f"{numbers['VP']}, and the conclusion.]", italic=True)
    doc.h1("2  Results")
    doc.table(["Activity", "Document", "Tests / checks", "Passed", "Failed", "Deviations"], [
        ["Installation Qualification", numbers["IQ"], "", "", "", ""],
        ["Operational Qualification", numbers["OQ"], str(sum(len(a.tests) for a in OQ)), "", "", ""],
        ["Performance Qualification", numbers["PQ"], str(len(PQ)), "", "", ""],
        ["Traceability", numbers["TM"], f"{len(REQUIREMENTS)} requirements", "", "", ""],
    ], [5, 3.5, 2.8, 1.8, 1.8, 2.1])
    doc.h1("3  Deviations")
    doc.table(["No.", "Test", "Description", "Impact / risk", "Corrective action", "Status"], [],
              [1.5, 2, 5, 3, 3.5, 2], blank_rows=4)
    doc.h1("4  Outstanding actions")
    doc.table(["Action", "Owner", "Due"], [], [11, 3, 3], blank_rows=3)
    doc.h1("5  Conclusion and release")
    doc.p("☐ The system is validated and released for GxP use.")
    doc.p("☐ The system is released with the restrictions listed above.")
    doc.p("☐ The system is not released.")
    return doc


# ---------------------------------------------------------------------------
# putting it together
# ---------------------------------------------------------------------------


@dataclass
class PackageResult:
    folder: Path
    files: list[Path]
    checks: list[Check]
    fingerprint: str

    @property
    def checks_passed(self) -> bool:
        return all(c.passed for c in self.checks)


def generate(db: Database, actor: SessionUser, destination: Path | str, prefix: str = "VAL-AURUM") -> PackageResult:
    """Write the complete package into a new dated folder under ``destination``."""
    require_any(actor, Permission.MANAGE_SETTINGS, Permission.APPROVE_RECORDS)
    stamp = datetime.now().astimezone()
    folder = Path(destination) / f"Aurum LIMS validation package {stamp:%Y-%m-%d %H%M}"
    folder.mkdir(parents=True, exist_ok=False)
    facts = installation_facts(db)
    with db.session() as session:
        lab = lab_info.get_lab_info(session).lab_name or "Laboratory name not set"
        security = get_security_settings(session)
        checks = installation_checks(db, session)
    manifest = program_manifest()
    fingerprint = manifest_fingerprint(manifest)
    numbers = _numbers(prefix)
    builders = {"VP": lambda: validation_plan(lab, numbers), "URS": lambda: urs(lab, numbers),
                "IQ": lambda: iq(lab, numbers, facts, checks, manifest, security), "OQ": lambda: oq(lab, numbers),
                "PQ": lambda: pq(lab, numbers), "TM": lambda: traceability(lab, numbers),
                "VSR": lambda: summary_report(lab, numbers)}
    files = [builders[code]().save(folder, prefix_) for prefix_, _, code in DOCUMENTS]
    with (folder / "program_file_manifest.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["File", "Size (bytes)", "SHA-256"])
        writer.writerows([e.path, e.size, e.sha256] for e in manifest)
        writer.writerow([])
        writer.writerow(["Overall program fingerprint", "", fingerprint])
    with (folder / "traceability_matrix.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Requirement", "Description", "GxP risk", "OQ tests", "PQ tests"])
        writer.writerows(traceability_rows())
    files += [folder / "program_file_manifest.csv", folder / "traceability_matrix.csv"]
    with db.session() as session:
        audit.record(session, AuditAction.VALIDATION_PACKAGE_CREATED, actor=actor, entity_type="validation",
                     entity_id=folder.name,
                     new={"folder": str(folder), "program_fingerprint": fingerprint,
                          "installation_checks": {c.name: "Pass" if c.passed else f"FAIL ({c.observed})"
                                                  for c in checks}})
    return PackageResult(folder, files, checks, fingerprint)


def last_package(session) -> AuditEntry | None:
    return session.scalar(select(AuditEntry).where(AuditEntry.action == AuditAction.VALIDATION_PACKAGE_CREATED)
                          .order_by(AuditEntry.id.desc()).limit(1))
